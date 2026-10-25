"""BRSR KPI mapping workbook contract tests (download + sheet/content validation)."""

import io
import os
import re
from typing import Dict, List

import pytest
import requests
from openpyxl import load_workbook


def _read_env_value(file_path: str, key: str) -> str:
    if not os.path.exists(file_path):
        return ""
    with open(file_path, "r", encoding="utf-8") as env_file:
        for raw_line in env_file:
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            if k.strip() == key:
                return v.strip().strip('"').strip("'")
    return ""


BASE_URL = os.environ.get("REACT_APP_BACKEND_URL") or _read_env_value("/app/frontend/.env", "REACT_APP_BACKEND_URL")


def _read_test_credentials() -> Dict[str, str]:
    with open("/app/memory/test_credentials.md", "r", encoding="utf-8") as credentials_file:
        content = credentials_file.read()
    email_match = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", content)
    password_match = re.search(r"Password\*\*:\s*(\S+)", content)
    assert email_match and password_match, "Test credentials are missing"
    return {"email": email_match.group(0), "password": password_match.group(1)}


@pytest.fixture(scope="session")
def api_session() -> requests.Session:
    session = requests.Session()
    session.headers.update({"Content-Type": "application/json"})
    return session


@pytest.fixture(scope="session")
def auth_headers(api_session: requests.Session) -> Dict[str, str]:
    # Module: Auth bootstrap for workbook download tests
    assert BASE_URL, "Missing REACT_APP_BACKEND_URL"
    credentials = _read_test_credentials()
    response = api_session.post(
        f"{BASE_URL}/api/auth/login",
        json=credentials,
        timeout=20,
    )
    assert response.status_code == 200, f"Login failed: {response.status_code}"
    token = response.json().get("access_token")
    assert token and isinstance(token, str)
    return {"Authorization": f"Bearer {token}"}


def _sheet_rows(sheet) -> List[List[str]]:
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values:
            continue
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append(["" if v is None else str(v) for v in values])
    return rows


def test_kpi_mappings_download_contract_and_filename(api_session: requests.Session, auth_headers: Dict[str, str]):
    # Feature: authenticated workbook download endpoint contract
    response = api_session.get(
        f"{BASE_URL}/api/brsr-report/kpi-mappings.xlsx",
        headers=auth_headers,
        timeout=60,
    )
    assert response.status_code == 200
    assert response.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    content_disposition = response.headers.get("content-disposition", "")
    assert "kpi_definitions_brsr_mappings.xlsx" in content_disposition
    assert len(response.content) > 1000


def test_kpi_mappings_workbook_sheets_and_columns(api_session: requests.Session, auth_headers: Dict[str, str]):
    # Feature: workbook sheet names and mapping sheet schema integrity
    response = api_session.get(
        f"{BASE_URL}/api/brsr-report/kpi-mappings.xlsx",
        headers=auth_headers,
        timeout=60,
    )
    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    assert workbook.sheetnames == ["Summary", "BRSR KPI Mappings", "Needs Your Response"]

    mapping_sheet = workbook["BRSR KPI Mappings"]
    headers = [cell.value for cell in mapping_sheet[1]]
    assert headers == [
        "KPI Name",
        "KPI Code",
        "Section",
        "Category",
        "Subcategory",
        "BRSR Reference",
        "BRSR Question Key",
        "Answer Path",
        "Aggregation Rule",
        "Formula Definition",
        "Imported From",
        "Mapping Status",
        "Needs Your Response",
        "Response Needed For",
    ]

    mapping_rows = _sheet_rows(mapping_sheet)
    assert len(mapping_rows) > 0


def test_needs_response_sheet_contains_only_flagged_rows(api_session: requests.Session, auth_headers: Dict[str, str]):
    # Feature: review sheet is a filtered subset containing only flagged mapping rows
    response = api_session.get(
        f"{BASE_URL}/api/brsr-report/kpi-mappings.xlsx",
        headers=auth_headers,
        timeout=60,
    )
    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    mapping_rows = _sheet_rows(workbook["BRSR KPI Mappings"])
    review_rows = _sheet_rows(workbook["Needs Your Response"])

    flagged_from_mapping = [row for row in mapping_rows if len(row) > 12 and row[12] == "YES"]
    assert len(review_rows) == len(flagged_from_mapping)
    assert all(len(row) > 12 and row[12] == "YES" for row in review_rows)
