"""Formula-driven GHG emissions summary workbook export."""
from collections import defaultdict
from io import BytesIO
import re
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SCOPE3_CATEGORIES = [
    "C1 - Purchased Goods and Services",
    "C2 - Capital Goods",
    "C3 - Fuel- and Energy-Related Activities",
    "C4 - Upstream Transportation and Distribution",
    "C5 - Waste Generated in Operations",
    "C6 - Business Travel",
    "C7 - Employee Commuting",
    "C8 - Upstream Leased Assets",
    "C9 - Downstream Transportation and Distribution",
    "C10 - Processing of Sold Products",
    "C11 - Use of Sold Products",
    "C12 - End-of-Life Treatment of Sold Products",
    "C13 - Downstream Leased Assets",
    "C14 - Franchises",
    "C15 - Investments",
]

THIN_GREY = Side(style="thin", color="CBD5E1")
TABLE_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)
TITLE_FILL = PatternFill("solid", fgColor="0F3D35")
SECTION_FILL = PatternFill("solid", fgColor="E6F1EE")
HEADER_FILL = PatternFill("solid", fgColor="F1F5F9")
TOTAL_FILL = PatternFill("solid", fgColor="EAF5EF")
NUMBER_FORMAT = '#,##0.00;[Red]-#,##0.00;'-''


def _emission_value(record: Dict[str, Any]) -> float:
    for field in ("total_emissions", "calculated_co2e", "co2e_emissions"):
        try:
            return float(record.get(field) or 0)
        except (TypeError, ValueError):
            continue
    return 0.0


def _scope_key(record: Dict[str, Any]) -> str:
    return re.sub(r"[^a-z0-9]", "", str(record.get("scope") or "").lower())


def _category_key(record: Dict[str, Any]) -> str:
    return re.sub(r"[^a-z0-9]", "", str(record.get("category") or "").lower())


def _scope1_bucket(record: Dict[str, Any]) -> str | None:
    category = _category_key(record)
    if "stationary" in category:
        return "Stationary Combustion"
    if "mobile" in category:
        return "Mobile Combustion"
    if "fugitive" in category:
        return "Fugitive Emissions"
    return None


def _scope2_bucket(record: Dict[str, Any]) -> str | None:
    category = _category_key(record)
    if any(token in category for token in ("heat", "steam")):
        return "Purchased Heat/Steam"
    if any(token in category for token in ("electric", "energy", "power")):
        return "Purchased Electricity"
    return None


def _scope3_bucket(record: Dict[str, Any]) -> str | None:
    category = str(record.get("category") or "").lower()
    code_match = re.search(r"(?:^|[^a-z0-9])c(?:ategory)?[_\s-]?(1[0-5]|[1-9])(?:$|[^0-9])", category)
    if code_match:
        return f"C{code_match.group(1)}"
    normalized = _category_key(record)
    aliases = {
        "purchasedgoodsandservices": "C1", "capitalgoods": "C2", "fuelandenergyrelatedactivities": "C3",
        "upstreamtransportationanddistribution": "C4", "wastegeneratedinoperations": "C5",
        "businesstravel": "C6", "employeecommuting": "C7", "upstreamleasedassets": "C8",
        "downstreamtransportationanddistribution": "C9", "processingofsoldproducts": "C10",
        "useofsoldproducts": "C11", "endoflifetreatmentofsoldproducts": "C12",
        "downstreamleasedassets": "C13", "franchises": "C14", "investments": "C15",
    }
    return aliases.get(normalized)


def _summary_data(records: Iterable[Dict[str, Any]], sinks: Iterable[Dict[str, Any]]) -> Dict[str, Any]:
    scope1, scope2, scope3 = defaultdict(float), defaultdict(float), defaultdict(float)
    biogenic = defaultdict(float)
    facilities = defaultdict(lambda: {"scope1": 0.0, "scope2": 0.0, "scope3": 0.0, "biogenic": 0.0, "sinks": 0.0})
    for record in records:
        amount = _emission_value(record)
        facility_id = record.get("facility_id")
        scope = _scope_key(record)
        if scope == "scope1":
            facilities[facility_id]["scope1"] += amount
            bucket = _scope1_bucket(record)
            if bucket:
                scope1[bucket] += amount
        elif scope == "scope2":
            facilities[facility_id]["scope2"] += amount
            bucket = _scope2_bucket(record)
            if bucket:
                scope2[bucket] += amount
        elif scope == "scope3":
            facilities[facility_id]["scope3"] += amount
            bucket = _scope3_bucket(record)
            if bucket:
                scope3[bucket] += amount
        elif scope == "biogenic":
            facilities[facility_id]["biogenic"] += amount
            selection = str(record.get("biogenic_scope_selection") or "").lower()
            biogenic["Indirect Biogenic" if selection in {"scope2", "indirect"} else "Direct Biogenic"] += amount
    for sink in sinks:
        facility_id = sink.get("facility_id")
        try:
            amount = float(sink.get("total_emissions_reduced") or 0) * float(sink.get("_proportion") or 1)
        except (TypeError, ValueError):
            amount = 0.0
        facilities[facility_id]["sinks"] += amount
    return {"scope1": scope1, "scope2": scope2, "scope3": scope3, "biogenic": biogenic, "facilities": facilities}


def _style_title(cell) -> None:
    cell.font = Font(bold=True, color="FFFFFF", size=14)
    cell.fill = TITLE_FILL


def _style_section(cell) -> None:
    cell.font = Font(bold=True, color="0F3D35", size=11)
    cell.fill = SECTION_FILL


def _style_header(cells) -> None:
    for cell in cells:
        cell.font = Font(bold=True, color="1E293B")
        cell.fill = HEADER_FILL
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(vertical="center")


def _style_total(cells) -> None:
    for cell in cells:
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL
        cell.border = TABLE_BORDER


def _write_detail_section(ws, start_row: int, title: str, rows: List[str], values: Dict[str, float]) -> tuple[int, int]:
    ws.cell(start_row, 1, title)
    _style_section(ws.cell(start_row, 1))
    header_row = start_row + 1
    ws.cell(header_row, 1, "Category")
    ws.cell(header_row, 2, "tCO2e")
    _style_header(ws[header_row][:2])
    first_data_row = header_row + 1
    for offset, label in enumerate(rows):
        row = first_data_row + offset
        ws.cell(row, 1, label)
        value_cell = ws.cell(row, 2, values[label] if label in values else "-")
        value_cell.number_format = NUMBER_FORMAT
        for cell in ws[row][:2]:
            cell.border = TABLE_BORDER
    total_row = first_data_row + len(rows)
    ws.cell(total_row, 1, f"Total {title}")
    total_cell = ws.cell(total_row, 2, f"=SUM(B{first_data_row}:B{total_row - 1})")
    total_cell.number_format = NUMBER_FORMAT
    _style_total(ws[total_row][:2])
    return total_row, total_row + 2


def build_ghg_summary_excel(
    organization: Dict[str, Any],
    facilities: List[Dict[str, Any]],
    records: List[Dict[str, Any]],
    sinks: List[Dict[str, Any]],
    reporting_period_start: str,
    reporting_period_end: str,
) -> bytes:
    """Create a two-sheet, formula-driven GHG emissions summary workbook."""
    data = _summary_data(records, sinks)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "GHG Emissions Summary"
    facility_sheet = workbook.create_sheet("Facility-Level Emissions")

    summary["A1"] = "GHG Emissions Summary"
    _style_title(summary["A1"])
    summary["A3"], summary["B3"] = "Organization Name", organization.get("name") or ""
    summary["D3"], summary["E3"] = "Reporting Year", f"{reporting_period_start[:4]}–{reporting_period_end[:4]}"
    summary["A4"], summary["B4"] = "Reporting Period", f"{reporting_period_start} to {reporting_period_end}"
    summary["D4"], summary["E4"] = "Number of Facilities", len(facilities)
    summary["A5"], summary["B5"] = "Reporting Boundary", organization.get("reporting_boundary") or organization.get("org_boundaries_approach") or ""
    summary["D5"], summary["E5"] = "Unit", "tCO2e"
    for row in range(3, 6):
        for col in (1, 2, 4, 5):
            summary.cell(row, col).border = TABLE_BORDER
        summary.cell(row, 1).font = summary.cell(row, 4).font = Font(bold=True)

    scope1_total_row, next_row = _write_detail_section(summary, 8, "Scope 1 Emissions", ["Stationary Combustion", "Mobile Combustion", "Fugitive Emissions"], data["scope1"])
    scope2_total_row, next_row = _write_detail_section(summary, next_row, "Scope 2 Emissions", ["Purchased Electricity", "Purchased Heat/Steam"], data["scope2"])
    scope3_total_row, next_row = _write_detail_section(summary, next_row, "Scope 3 Emissions", SCOPE3_CATEGORIES, {label: data["scope3"].get(label.split(" - ", 1)[0], 0) for label in SCOPE3_CATEGORIES if data["scope3"].get(label.split(" - ", 1)[0]) is not None})
    biogenic_total_row, next_row = _write_detail_section(summary, next_row, "Biogenic Emissions", ["Direct Biogenic", "Indirect Biogenic"], data["biogenic"])

    summary.cell(next_row, 1, "Total Emissions")
    _style_section(summary.cell(next_row, 1))
    summary.cell(next_row + 1, 1, "Metric")
    summary.cell(next_row + 1, 2, "tCO2e")
    _style_header(summary[next_row + 1][:2])
    summary_rows = [
        ("Scope 1 Emissions", f"=B{scope1_total_row}"),
        ("Scope 2 Emissions", f"=B{scope2_total_row}"),
        ("Scope 3 Emissions", f"=B{scope3_total_row}"),
        ("Total GHG Emissions", f"=SUM(B{next_row + 2}:B{next_row + 4})"),
        ("Total Biogenic Emissions", f"=B{biogenic_total_row}"),
        ("Sinks", None),
        ("Net Emissions", f"=B{next_row + 5}-B{next_row + 7}"),
    ]
    for offset, (label, formula) in enumerate(summary_rows, start=next_row + 2):
        summary.cell(offset, 1, label)
        value_cell = summary.cell(offset, 2, formula or "")
        value_cell.number_format = NUMBER_FORMAT
        for cell in summary[offset][:2]:
            cell.border = TABLE_BORDER
    facility_total_row = 6 + len(facilities)
    summary.cell(next_row + 7, 2, f"='Facility-Level Emissions'!F{facility_total_row}")
    _style_total(summary[next_row + 5][:2])
    _style_total(summary[next_row + 8][:2])

    facility_sheet["A1"] = "Facility-Level Emissions"
    _style_title(facility_sheet["A1"])
    facility_sheet["A3"] = "Reporting Period"
    facility_sheet["B3"] = f"{reporting_period_start} to {reporting_period_end}"
    facility_sheet["A3"].font = Font(bold=True)
    facility_sheet["A3"].border = facility_sheet["B3"].border = TABLE_BORDER
    headers = ["Facility", "Scope 1 (tCO2e)", "Scope 2 (tCO2e)", "Scope 3 (tCO2e)", "Biogenic (tCO2e)", "Sinks (tCO2e)", "Total GHG Emissions (tCO2e)", "Net Emissions (tCO2e)"]
    for column, label in enumerate(headers, start=1):
        facility_sheet.cell(5, column, label)
    _style_header(facility_sheet[5][:8])
    for index, facility in enumerate(facilities, start=6):
        values = data["facilities"].get(facility.get("id"), {})
        facility_sheet.cell(index, 1, facility.get("name") or "Facility")
        for column, key in enumerate(("scope1", "scope2", "scope3", "biogenic", "sinks"), start=2):
            cell = facility_sheet.cell(index, column, values[key] if key in values else "-")
            cell.number_format = NUMBER_FORMAT
        facility_sheet.cell(index, 7, f"=SUM(B{index}:D{index})").number_format = NUMBER_FORMAT
        facility_sheet.cell(index, 8, f"=G{index}-F{index}").number_format = NUMBER_FORMAT
        for cell in facility_sheet[index][:8]:
            cell.border = TABLE_BORDER
    facility_sheet.cell(facility_total_row, 1, "Total")
    for column in range(2, 9):
        cell = facility_sheet.cell(facility_total_row, column, f"=SUM({get_column_letter(column)}6:{get_column_letter(column)}{facility_total_row - 1})")
        cell.number_format = NUMBER_FORMAT
    _style_total(facility_sheet[facility_total_row][:8])

    for sheet in (summary, facility_sheet):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A6" if sheet == facility_sheet else "A8"
        for row in sheet.iter_rows():
            for cell in row:
                if cell.column > 1 and isinstance(cell.value, (int, float)):
                    cell.number_format = NUMBER_FORMAT
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 4
    summary.column_dimensions["D"].width = 22
    summary.column_dimensions["E"].width = 24
    for column, width in {"A": 30, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 25, "H": 22}.items():
        facility_sheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()