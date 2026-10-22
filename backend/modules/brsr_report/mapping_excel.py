import io
import hashlib
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shared.database.mongo import db


HEADERS = [
    "KPI Name",
    "KPI Code",
    "Section",
    "Category",
    "Subcategory",
    "BRSR Reference",
    "BRSR Question Key",
    "Answer Path",
    "Aggregation Rule",
    "Formula Definition",
    "Imported From",
    "Mapping Status",
    "Needs Your Response",
    "Response Needed For",
]


def _has_brsr_reference(kpi: Dict[str, Any]) -> bool:
    if kpi.get("brsr_reference"):
        return True
    return any(
        str(reference.get("framework", "")).upper() == "BRSR"
        for reference in (kpi.get("framework_references") or [])
        if isinstance(reference, dict)
    )


def _review_reason(mapping: Dict[str, Any]) -> str:
    missing = []
    if not mapping:
        return "No BRSR mapping record"
    if not mapping.get("question_key"):
        missing.append("question key")
    if not mapping.get("answer_path"):
        missing.append("answer path")
    if mapping.get("mapping_status") != "exact_destination" and not missing:
        missing.append("exact destination confirmation")
    return ", ".join(missing)


def _question_answer_paths(config: Dict[str, Any], mapped_paths: List[str]) -> List[str]:
    paths = {path for path in mapped_paths if path}
    table_config = config.get("table_config") or {}
    for row_key in ("fixed_rows", "rows"):
        for row in table_config.get(row_key) or []:
            if isinstance(row, dict) and row.get("key"):
                paths.add(f"rows.{row['key']}")
    if not paths:
        paths.add("value")
    return sorted(paths)


def _range_name(question_key: str) -> str:
    digest = hashlib.sha1(question_key.encode("utf-8")).hexdigest()[:12]
    return f"BRSR_Answers_{digest}"


def _add_mapping_dropdowns(workbook: Workbook, sheet, last_row: int) -> None:
    if last_row < 2:
        return
    question_validation = DataValidation(
        type="list",
        formula1="=BRSR_Question_Keys",
        allow_blank=True,
        errorStyle="stop",
        errorTitle="Invalid BRSR question",
        error="Choose a BRSR Question Key from the dropdown.",
        promptTitle="BRSR Question Key",
        prompt="Select a configured BRSR question.",
        showInputMessage=True,
        showErrorMessage=True,
    )
    answer_validation = DataValidation(
        type="list",
        formula1='=INDIRECT(IFERROR(VLOOKUP($G2,\'BRSR Options\'!$A:$C,3,FALSE),"BRSR_No_Options"))',
        allow_blank=True,
        errorStyle="stop",
        errorTitle="Invalid answer path",
        error="Choose an Answer Path available for the selected BRSR Question Key.",
        promptTitle="BRSR Answer Path",
        prompt="Select the question first, then choose its answer path.",
        showInputMessage=True,
        showErrorMessage=True,
    )
    sheet.add_data_validation(question_validation)
    sheet.add_data_validation(answer_validation)
    question_validation.add(f"G2:G{last_row}")
    answer_validation.add(f"H2:H{last_row}")


async def build_kpi_brsr_mapping_excel() -> bytes:
    kpis = await db.esg_kpi_definitions.find(
        {"status": {"$ne": "archived"}},
        {
            "_id": 0,
            "id": 1,
            "metric_name": 1,
            "metric_code": 1,
            "section": 1,
            "category_name": 1,
            "subcategory": 1,
            "brsr_reference": 1,
            "framework_references": 1,
        },
    ).sort([("section", 1), ("category_name", 1), ("subcategory", 1), ("metric_name", 1)]).to_list(10000)
    kpis = [kpi for kpi in kpis if _has_brsr_reference(kpi)]

    mappings = await db.kpi_framework_mappings.find(
        {"framework": "BRSR"}, {"_id": 0}
    ).to_list(10000)
    mapped_paths_by_question: Dict[str, List[str]] = {}
    for mapping in mappings:
        if mapping.get("question_key") and mapping.get("answer_path"):
            mapped_paths_by_question.setdefault(mapping["question_key"], []).append(mapping["answer_path"])

    question_configs = await db.esg_question_configs.find(
        {"$or": [{"frameworks": "BRSR"}, {"question_key": {"$regex": "^p[0-9]+_", "$options": "i"}}]},
        {"_id": 0, "question_key": 1, "question": 1, "table_config": 1, "order": 1},
    ).sort([("question_key", 1)]).to_list(10000)
    question_options = []
    for config in question_configs:
        question_key = config.get("question_key")
        if not question_key:
            continue
        question_options.append({
            "question_key": question_key,
            "question": config.get("question") or question_key,
            "answer_paths": _question_answer_paths(config, mapped_paths_by_question.get(question_key, [])),
            "range_name": _range_name(question_key),
        })
    mapping_by_kpi: Dict[str, List[Dict[str, Any]]] = {}
    for mapping in mappings:
        source_ids = mapping.get("source_kpi_ids") or [mapping.get("kpi_id")]
        for kpi_id in source_ids:
            if kpi_id:
                mapping_by_kpi.setdefault(kpi_id, []).append(mapping)

    rows = []
    review_rows = []
    for kpi in kpis:
        kpi_mappings = mapping_by_kpi.get(kpi["id"]) or [{}]
        for mapping in kpi_mappings:
            reason = _review_reason(mapping)
            needs_response = bool(reason)
            row = [
                kpi.get("metric_name"),
                kpi.get("metric_code"),
                kpi.get("section"),
                kpi.get("category_name"),
                kpi.get("subcategory"),
                mapping.get("reference") or kpi.get("brsr_reference"),
                mapping.get("question_key"),
                mapping.get("answer_path"),
                mapping.get("aggregation_rule"),
                mapping.get("formula_definition"),
                mapping.get("source_tag") or " → ".join(filter(None, [kpi.get("category_name"), kpi.get("subcategory")])),
                mapping.get("mapping_status") or "unmapped",
                "YES" if needs_response else "NO",
                reason,
            ]
            rows.append(row)
            if needs_response:
                review_rows.append(row)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "BRSR KPI Mappings"
    review_sheet = workbook.create_sheet("Needs Your Response")
    options_sheet = workbook.create_sheet("BRSR Options")

    options_sheet.append(["BRSR Question Key", "Question", "Answer Path Range"])
    for row_number, option in enumerate(question_options, start=2):
        options_sheet.cell(row_number, 1, option["question_key"])
        options_sheet.cell(row_number, 2, option["question"])
        options_sheet.cell(row_number, 3, option["range_name"])
        answer_column = 3 + row_number - 1
        options_sheet.cell(1, answer_column, option["question_key"])
        for answer_row, answer_path in enumerate(option["answer_paths"], start=2):
            options_sheet.cell(answer_row, answer_column, answer_path)
        column_letter = get_column_letter(answer_column)
        workbook.defined_names.add(DefinedName(
            option["range_name"],
            attr_text=f"'BRSR Options'!${column_letter}$2:${column_letter}${len(option['answer_paths']) + 1}",
        ))

    question_last_row = max(2, len(question_options) + 1)
    workbook.defined_names.add(DefinedName(
        "BRSR_Question_Keys",
        attr_text=f"'BRSR Options'!$A$2:$A${question_last_row}",
    ))
    no_options_row = question_last_row + 2
    options_sheet.cell(no_options_row, 1, "")
    workbook.defined_names.add(DefinedName(
        "BRSR_No_Options",
        attr_text=f"'BRSR Options'!$A${no_options_row}",
    ))

    header_fill = PatternFill("solid", fgColor="14532D")
    review_fill = PatternFill("solid", fgColor="FEF3C7")
    exact_fill = PatternFill("solid", fgColor="DCFCE7")
    thin_border = Border(bottom=Side(style="thin", color="D6D3D1"))

    for sheet, data_rows in ((worksheet, rows), (review_sheet, review_rows)):
        sheet.append(HEADERS)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:N{max(1, len(data_rows) + 1)}"

        for row_index, row in enumerate(data_rows, start=2):
            sheet.append(row)
            for cell in sheet[row_index]:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row[12] == "YES":
                for cell in sheet[row_index]:
                    cell.fill = review_fill
                sheet.cell(row_index, 13).font = Font(color="92400E", bold=True)
            else:
                sheet.cell(row_index, 12).fill = exact_fill
                sheet.cell(row_index, 13).fill = exact_fill
                sheet.cell(row_index, 13).font = Font(color="166534", bold=True)

        widths = [34, 22, 14, 24, 28, 18, 24, 30, 20, 40, 34, 24, 20, 34]
        for column_index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        sheet.row_dimensions[1].height = 34
        sheet["G1"].comment = Comment(
            "Choose from configured BRSR questions. One KPI may appear on multiple mapping rows and point to multiple questions.",
            "SustainRepo",
        )
        sheet["H1"].comment = Comment(
            "This dropdown changes based on the selected BRSR Question Key.",
            "SustainRepo",
        )
        _add_mapping_dropdowns(workbook, sheet, len(data_rows) + 1)

    summary = workbook.create_sheet("Summary", 0)
    summary.append(["BRSR KPI Mapping Review", "Count"])
    summary.append(["KPI mapping rows", len(rows)])
    summary.append(["Ready / exact destination", len(rows) - len(review_rows)])
    summary.append(["Needs your response", len(review_rows)])
    summary.append(["Review rule", "Rows need your response when question key, answer path, or exact destination confirmation is missing."])
    summary.append(["Question Key options", "Use the dropdown in BRSR Question Key. Options come from the configured BRSR questionnaire catalogue."])
    summary.append(["Answer Path options", "Select a Question Key first; the Answer Path dropdown then shows only that question's configured paths."])
    summary.append(["Multiple destinations", "One KPI can map to multiple destinations by using multiple mapping rows. This workbook contains current rows only, as selected."])
    summary.append(["Persistence", "Workbook selections remain in the downloaded file; they are not applied to the application until an import workflow is implemented."])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 95
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    summary.freeze_panes = "A2"
    options_sheet.sheet_state = "hidden"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()