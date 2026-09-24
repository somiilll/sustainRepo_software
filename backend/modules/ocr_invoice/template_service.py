"""Reference-format OCR activity spreadsheet template."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


LEDGER_HEADERS = [
    "Facility (Mandatory)",
    "Reporting Period (Mandatory)",
    "Invoice Number (Mandatory)",
    "Vendor Name",
    "Item Description (Mandatory)",
    "Cost (Mandatory, Required if entry is related to travel or transport)",
    "Quantity",
    "Units",
    "Distance Travelled (Required if entry is related to travel or transport)",
    "No. of Passengers Travelled (Required if entry is related to travel or transport)",
    "Number of Rooms (Required if hotel stay)",
    "Number of Nights (Required if hotel stay)",
    "From Location (Required if entry is related to travel or transport)",
    "To Location (Required if entry is related to travel or transport)",
    "Notes",
]
GUIDE_ROWS = [
    ("Facility", "The facility, plant, office, or site where this activity occurred.", "Mandatory. Select a value from the dropdown."),
    ("Reporting Period", "The month and year or financial year for the activity.", "Mandatory. Example: April 2026 or FY 2025-26."),
    ("Invoice Number", "The invoice, bill, or internal reference number.", "Mandatory when available."),
    ("Vendor Name", "The supplier, service provider, or travel provider.", "Use the invoice vendor name where available."),
    ("Item Description", "A clear description of the purchased item, fuel, transport, travel, accommodation, or service.", "Mandatory. Helps OCR classify the activity."),
    ("Cost", "The spend amount for the invoice line or activity.", "Required for travel and transport entries."),
    ("Quantity / Units", "Measured activity quantity and its unit, such as L, kg, tonnes, kWh, m3, or km.", "Use when a physical quantity is available."),
    ("Travel fields", "Distance, passenger count, origin, and destination for passenger or goods travel.", "Required for travel or transport entries where applicable."),
    ("Hotel fields", "Number of rooms and number of nights for accommodation.", "Required for hotel stays."),
    ("Notes", "Additional context that can help classify the activity.", "Optional."),
]
UNIT_ROWS = [
    ("L", "Litres"),
    ("kg", "Kilograms"),
    ("tonnes", "Tonnes"),
    ("kWh", "Kilowatt-hour"),
    ("m3", "Cubic meters"),
    ("km", "Kilometers"),
    ("No.", "Number"),
    ("Nights", "Hotel nights"),
    ("Rooms", "Hotel rooms"),
]

THIN_BLACK = Side(style="thin", color="000000")
THIN_GREY = Side(style="thin", color="B7B7B7")
HEADER_BORDER = Border(left=THIN_BLACK, right=THIN_BLACK, top=THIN_BLACK, bottom=THIN_BLACK)
GRID_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)
TITLE_FILL = PatternFill("solid", fgColor="ADD8E6")
MANDATORY_FILL = PatternFill("solid", fgColor="FF0000")
CONDITIONAL_FILL = PatternFill("solid", fgColor="FFA500")
OPTIONAL_FILL = PatternFill("solid", fgColor="FFFF00")
GUIDE_TITLE_FILL = PatternFill("solid", fgColor="D3D3D3")
GUIDE_HEADER_FILL = PatternFill("solid", fgColor="ADD8E6")


def _set_cell_border_and_alignment(cell, border, horizontal="left"):
    cell.border = border
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)


async def generate_ocr_template(database, organization_id: str) -> BytesIO:
    """Build the OCR Activity Ledger workbook in the approved reference format."""
    facilities = await database.facilities.find(
        {
            "organization_id": organization_id,
            "is_deleted": {"$ne": True},
            "is_active": {"$ne": False},
        },
        {"_id": 0, "name": 1},
    ).sort("name", 1).to_list(1000)
    facility_names = [str(facility["name"]).strip() for facility in facilities if facility.get("name")]

    workbook = Workbook()
    ledger = workbook.active
    ledger.title = "OCR Activity Ledger"
    ledger.sheet_view.showGridLines = False
    ledger.merge_cells("A1:O1")
    ledger["A1"] = "OCR ACTIVITY LEDGER"
    ledger["A1"].font = Font(name="Calibri", size=18, bold=True)
    ledger["A1"].fill = TITLE_FILL
    _set_cell_border_and_alignment(ledger["A1"], HEADER_BORDER, "center")
    ledger.row_dimensions[1].height = 25
    ledger.row_dimensions[2].height = 30
    ledger.freeze_panes = "B3"
    ledger.auto_filter.ref = "A2:O30"

    widths = [18, 18, 18, 25, 40, 15, 15, 12, 18, 22, 18, 18, 25, 25, 35]
    for column, width in enumerate(widths, start=1):
        ledger.column_dimensions[get_column_letter(column)].width = width

    for column, header in enumerate(LEDGER_HEADERS, start=1):
        cell = ledger.cell(row=2, column=column, value=header)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.fill = MANDATORY_FILL if column <= 5 else CONDITIONAL_FILL if column <= 14 else OPTIONAL_FILL
        _set_cell_border_and_alignment(cell, HEADER_BORDER, "center")

    for row in range(3, 31):
        ledger.row_dimensions[row].height = 20
        for column in range(1, 16):
            cell = ledger.cell(row=row, column=column)
            cell.font = Font(name="Calibri", size=11)
            cell.protection = Protection(locked=False)
            _set_cell_border_and_alignment(cell, GRID_BORDER, "right" if 6 <= column <= 12 else "left")
        ledger.cell(row=row, column=6).number_format = '₹#,##0.00'
        ledger.cell(row=row, column=9).number_format = "0"
        for column in (10, 11, 12):
            ledger.cell(row=row, column=column).number_format = "0"

    ledger.conditional_formatting.add(
        "A3:O30",
        FormulaRule(formula=['$A3=""'], fill=PatternFill("solid", fgColor="DCDCDC")),
    )
    for column in (6, 9, 11):
        letter = get_column_letter(column)
        ledger.conditional_formatting.add(
            f"{letter}3:{letter}30",
            FormulaRule(formula=[f'AND(${letter}3<>"",$A3="")'], font=Font(color="FF0000")),
        )

    facility_codes = workbook.create_sheet("Formatting Codes")
    facility_codes.sheet_state = "hidden"
    facility_codes.append(["Facility Name", "Description (Internal Use)"])
    for facility_name in facility_names:
        facility_codes.append([facility_name, "Organization facility"])
    if not facility_names:
        facility_codes.append(["", "Organization facility"])
    for row in facility_codes.iter_rows():
        for cell in row:
            _set_cell_border_and_alignment(cell, HEADER_BORDER)

    unit_definitions = workbook.create_sheet("Unit Definitions")
    unit_definitions.sheet_state = "hidden"
    unit_definitions.append(["Unit", "Description (Internal Use)"])
    for unit, description in UNIT_ROWS:
        unit_definitions.append([unit, description])
    for row in unit_definitions.iter_rows():
        for cell in row:
            _set_cell_border_and_alignment(cell, HEADER_BORDER)

    if facility_names:
        facility_validation = DataValidation(
            type="list",
            formula1=f"='Formatting Codes'!$A$2:$A${len(facility_names) + 1}",
            allow_blank=True,
        )
        facility_validation.errorTitle = "Unknown facility"
        facility_validation.error = "Choose a facility from your organization list."
        ledger.add_data_validation(facility_validation)
        facility_validation.add("A3:A30")
    unit_validation = DataValidation(
        type="list",
        formula1=f"='Unit Definitions'!$A$2:$A${len(UNIT_ROWS) + 1}",
        allow_blank=True,
    )
    ledger.add_data_validation(unit_validation)
    unit_validation.add("H3:H30")

    guide = workbook.create_sheet("Guide", 1)
    guide.sheet_view.showGridLines = False
    guide.merge_cells("A1:C1")
    guide["A1"] = "OCR Activity Template Guide"
    guide["A1"].font = Font(name="Calibri", size=16, bold=True)
    guide["A1"].fill = GUIDE_TITLE_FILL
    _set_cell_border_and_alignment(guide["A1"], HEADER_BORDER, "center")
    guide.row_dimensions[1].height = 20
    for column, header in enumerate(["Field", "Description", "Notes / Guidance"], start=1):
        cell = guide.cell(row=2, column=column, value=header)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.fill = GUIDE_HEADER_FILL
        _set_cell_border_and_alignment(cell, HEADER_BORDER, "center")
    for row, guide_row in enumerate(GUIDE_ROWS, start=3):
        for column, value in enumerate(guide_row, start=1):
            cell = guide.cell(row=row, column=column, value=value)
            cell.font = Font(name="Calibri", size=11)
            _set_cell_border_and_alignment(cell, GRID_BORDER, "center" if column == 1 else "left")
        guide.row_dimensions[row].height = 34
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 60
    guide.column_dimensions["C"].width = 35

    ledger.protection.sheet = True
    ledger.protection.selectLockedCells = True
    ledger.protection.selectUnlockedCells = True
    guide.protection.sheet = True
    facility_codes.protection.sheet = True
    unit_definitions.protection.sheet = True

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output