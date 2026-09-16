"""Document rendering, structured extraction, and spreadsheet ingestion."""
from __future__ import annotations

import base64
import csv
import os
import re
from io import BytesIO
from pathlib import Path
from typing import Any

import fitz
import openpyxl
import xlrd
from PIL import Image

from .companion_rows import build_companion_rows
from .config import MAX_PDF_PAGES
from .llm_gateway import OcrLlmGateway
from .normalization import (
    normalize_currency,
    normalize_date,
    normalize_period,
    normalize_unit,
    repair_extraction_json,
)
from .reconciliation import finalize_spreadsheet_item, prepare_item, prepare_spreadsheet_item, reconcile_invoice_items
from .taxonomy_service import classify_item


EXTRACTION_SYSTEM_PROMPT = """You are an expert ESG emissions accounting data extraction assistant.
Extract transaction details and PHYSICAL ACTIVITY EVENTS from the provided invoice(s) into raw JSON format. Return ONLY raw JSON without markdown blocks or backticks.

UNIVERSAL PHYSICAL ACTIVITY EXTRACTION RULES:
1. MULTI-INVOICE / MULTI-PURCHASE-ORDER DOCUMENTS:
   - If the document contains multiple distinct invoices, purchase orders, or delivery manifests across different pages, return a JSON ARRAY of invoice objects: `[{...}, {...}]`.
   - If it is a single invoice with 1 or multiple items, return a JSON object `{...}` or an array with 1 object.
2. HIERARCHICAL LINE ITEM DE-DUPLICATION (CRITICAL):
   - In SAP/ERP purchase orders and service bills, DO NOT extract summary parent group headers if itemized leaf lines are present with individual quantities and rates.
   - Example: If a PO has Line 10 "REFILLING OF FIRE EXTINGUISHER Amount: 11075" and then Lines 1-4 list ABC 4KG, ABC 6KG, CO2 4.5KG, CO2 9KG whose amounts sum to 11075:
     * Extract ONLY the 4 leaf breakdown items!
     * DROP the parent summary header to prevent double counting.
3. Identify the PRIMARY PHYSICAL ACTIVITY ITEMS for each invoice (physical goods, raw materials, energy, waste, travel, freight).
4. Classify `material_nature` for each item:
   - "raw_material": Single homogeneous unprocessed/semi-processed material (e.g. copper wire rod, steel ingots, plastic granules, bulk chemicals).
   - "composite_product": Manufactured, assembled, or multi-material finished products (e.g. electrical appliances, electric motors, PCB assemblies, tools, valves, equipment, furniture).
   - "service": Non-physical services (e.g. legal, consulting, SaaS, software licenses, maintenance labor).
   - "energy": Fuels, electricity, steam, district heating/cooling.
   - "waste": Refuse, scrap, wastewater, recycling manifests.
   - "logistics": Freight, courier, transport services.
   - "travel": Air tickets, hotel stays, rail tickets, taxi rides.
5. If item is raw_material, specify `primary_material` (e.g. Copper, Steel, Aluminum, LDPE, Paper).
6. DO NOT extract itemized fee components (such as tipping fees, hauling rates, loading/handling charges, meter fees, fixed demand charges, fuel surcharges, late fees, documentation fees, or taxes) as separate primary activity items. Capture financial sums in `invoice_financials`.
7. PHYSICAL QUANTITY & PACKAGING NORMALIZATION:
   - For any fuel, gas, liquid, chemical, or raw material delivered in packaging, containers, or item counts (cylinders, drums, bottles, bags, barrels, cans, tankers, boxes, EA, Nos, Pcs):
     * Always calculate and report the TOTAL NET PHYSICAL MASS OR VOLUME in standard canonical units:
       - Solid / Mass: kg or tonnes (e.g., "11 EA of 4 KG ABC" -> quantity = 44.0, unit = "kg"; "5 EA of 4.5 KG CO2" -> quantity = 22.5, unit = "kg").
       - Liquid / Volume: Liters, Gallons, m3, kL.
       - Gas: kg (for CO2, LPG, refrigerants) or m3.
       - Energy: kWh, MWh.
       - Distance: km.
     * DO NOT use EA, Nos, Pcs, or container counts as the unit if capacity is specified.
     * If an item is a pure non-physical service (e.g. consulting, maintenance labor with no chemical/gas), report quantity = null, unit = null.
8. CONFIDENCE SCORE & CERTAINTY RULES:
   - Evaluate confidence based on visual clarity, explicit printing vs. inferences, and emissions calculation feasibility:
     * EXPLICIT DATA (HIGH CONFIDENCE: 85-98%):
       - If physical quantity AND explicit unit (e.g., "56 m3", "1000 Liters", "450 kWh") are clearly printed on the document, assign HIGH confidence (90-98%), even if cost is not broken out.
       - If total financial cost (e.g., "$5,000", "₹1,50,000") is clearly printed for purchased goods/services where physical quantity is not applicable, assign HIGH confidence (85-95%).
       - When all required fields are explicitly printed and legible, leave low_confidence_fields empty [].
     * MISSING / INFERRED UNIT (MEDIUM CONFIDENCE: 70-75%):
       - If a physical quantity number is present (e.g. "70" under Qty in UnE) but the UNIT OF MEASURE IS NOT EXPLICITLY PRINTED on the document and you had to infer/assume it from context (e.g. inferred "Liters" for diesel from forklift equipment context):
         - Assign a confidence score of 70-75% (Medium confidence).
         - Add "inferred unit" (or "missing unit" if null) to low_confidence_fields.
         - In additional_context, explain the assumption (e.g. "Unit was not printed on document; inferred Liters based on diesel fuel in forklift.").
     * CRITICAL DATA GAP (LOW CONFIDENCE: <= 40%):
       - If BOTH physical activity (quantity/distance) AND financial expenditure (cost/amount) are missing, score must be <= 40%, and add "missing qty & cost" to low_confidence_fields.
     * NON-CRITICAL AUXILIARY FIELDS (DO NOT LOWER SCORE):
       - DO NOT flag vendor_name, vendor_type, vendor_address, buyer details, or unit_price in low_confidence_fields. These auxiliary fields must not lower the score or trigger warning flags.
      * OCR / LEGIBILITY ISSUES:
        - If text is blurry, truncated, or digits are ambiguous, assign 50-65% and flag the specific field (e.g. "blurry quantity").
9. BUSINESS TRAVEL / PASSENGER TICKETS (Scope 3 Category 6):
   - When extracting tickets, boarding passes, itineraries, or travel bills:
     * Set `item_category_hint: "Travel"` and `material_nature: "travel"`.
     * Railways: Extract `travel_details.mode = "Train"`, `origin` (departure station/city), `destination` (arrival station/city), and `passenger_count` (int, default 1). Leave `class: null` (no cabin class for rail).
     * Road / Cabs / Taxis (Uber, Ola, local taxi, rental): Extract `travel_details.mode = "Taxi"` or `"Car"`, `origin`, `destination`, `passenger_count`, and NORMALIZE vehicle type strictly to DEFRA size brackets in `travel_details.vehicle_type`: "Small", "Medium", "Large", or "Average" (default to "Average" if unspecified).
     * Aviation / Flights: Extract `travel_details.mode = "Flight"`, `origin` (departure airport code or city), `destination` (arrival airport code or city), `class` ("Economy" | "Premium Economy" | "Business" | "First"), and `passenger_count` (int, default 1). DO NOT calculate point-to-point distance in km; leave `distance_km: null` if not explicitly printed on the ticket.
     * Hotels / Accommodation: Extract `travel_details.mode = "Hotel"`, `origin` (city/country of hotel), calculate total room nights as `quantity` = (number of rooms) * (number of nights), and set `unit: "room_nights"`.

JSON SCHEMA PER INVOICE:
{
    "invoice_number": "string or null",
    "date": "YYYY-MM-DD or null",
    "billing_period_start": "YYYY-MM-DD or null",
    "billing_period_end": "YYYY-MM-DD or null",
    "vendor_name": "string or null",
    "vendor_type": "Supplier/Vendor | Utility | Waste Management | Travel Agency | Freight/Logistics | Fuel Station | Refrigerant Supplier | Water Utility | Government | Other",
    "vendor_address": "string or null",
    "buyer_name": "Name of the client / billed-to entity or null",
    "buyer_address": "Address of the billed-to entity or null",
    "service_address": "Full address where service/goods were DELIVERED (the client's destination/plant/facility), NOT the vendor's billing/dispatch office or null",
    "invoice_financials": {
        "subtotal": <float or null>,
        "tax_amount": <float or null>,
        "additional_charges_amount": <float or null>,
        "grand_total": <float or null>,
        "currency": "3-letter code (e.g. INR, USD, EUR) or null"
    },
    "primary_activity_items": [
        {
            "item_description": "Name of the physical product, service, or activity",
            "item_description_english": "English translation or standard description",
            "item_category_hint": "Purchased Goods | Capital Goods | Fuel/Energy | Water | Waste | Travel | Freight/Logistics | Refrigerant | Employee Commute | Other",
            "material_nature": "raw_material | composite_product | service | energy | waste | logistics | travel",
            "primary_material": "Dominant raw material if raw_material (e.g. Copper, Steel, Aluminum, LDPE, Paper). null otherwise",
            "quantity": <float or null>,
            "unit": "kg | tonnes | Liters | Gallons | kWh | km | miles | nights | room_nights | trips | m3 | kL | null",
            "unit_price": <float or null>,
            "base_cost": <float or null>,
            "waste_details": {
                "waste_type": "string or null",
                "disposal_method": "Landfill | Recycling | Composting | Incineration | Wastewater | null"
            },
            "travel_details": {
                "mode": "Flight | Train | Taxi | Bus | Car Rental | Hotel | null",
                "class": "Economy | Premium Economy | Business | First | null",
                "vehicle_type": "Small | Medium | Large | Average | null",
                "passenger_count": <int or null>,
                "distance_km": <float or null>,
                "origin": "string or null",
                "destination": "string or null"
            },
            "freight_details": {
                "mode": "Road | Rail | Sea | Air | Courier | null",
                "weight_kg": <float or null>,
                "distance_km": <float or null>,
                "origin": "string or null",
                "destination": "string or null"
            },
            "additional_context": "any other relevant context from the invoice",
            "confidence_score": <int 0-100>,
            "low_confidence_fields": ["array of low confidence keys", "or empty"]
        }
    ]
}"""


EXTRACTION_SCHEMA_PROMPT = "Extract the physical activity events and financial breakdown into the requested JSON format."


HEADER_ALIASES = {
    "facility": ["facility_name", "facility", "plant_name", "plant", "site", "branch_name", "branch", "cost_center", "location"],
    "reporting_period": ["reporting_period", "reporting_month", "financial_year", "reporting_period_month_or_fy", "period", "billing_period", "fin_year", "fiscal_year", "fy", "month", "accounting_period", "reporting_year"],
    "invoice_number": ["invoice_number", "invoicenumber", "inv_no", "inv_num", "invoice_no", "bill_no", "bill_number", "doc_no", "document_no", "voucher_no"],
    "date": ["invoice_date", "invoicedate", "date", "bill_date", "doc_date", "posting_date"],
    "vendor_name": ["vendor_name", "vendor", "supplier_name", "supplier", "party_name", "party", "payee"],
    "item_description": ["item_description", "description", "item", "particulars", "material_description", "material_name", "product_name", "product", "activity", "service_description"],
    "quantity": ["quantity", "qty", "billed_qty", "volume", "net_weight", "quantity_used", "units"],
    "unit": ["unit", "uom", "unit_of_measure", "unit_of_quantity", "unit_of_quantity_used", "measure"],
    "total_cost": ["total_cost", "total_amount", "amount", "cost", "total", "net_amount", "spent_amount", "spent_amount_inr", "spend", "invoice_amount", "value"],
    "currency": ["currency", "curr"],
    "distance_km": ["distance_km", "distance", "dist_km", "dist", "travel_distance", "km", "kms", "route_distance"],
    "origin": ["origin", "from", "source", "departure", "pickup"],
    "destination": ["destination", "to", "arrival", "drop", "dest"],
    "primary_material": ["primary_material", "material", "material_type", "dominant_material"],
    "material_nature": ["material_nature", "nature", "item_nature"],
    "item_category_hint": ["item_category_hint", "category_hint", "activity_type"],
    "hsn_sac_code": ["hsn_sac_code", "hsn", "sac", "hsn_code", "sac_code"],
    "unit_price": ["unit_price", "rate", "price_per_unit"],
    "travel_mode": ["travel_mode", "mode_of_travel", "transport_mode"],
    "travel_class": ["travel_class", "class", "cabin_class", "cabin", "booking_class"],
    "vehicle_type": ["vehicle_type", "vehicle", "car_type", "cab_type"],
    "passenger_count": ["passenger_count", "passengers", "pax"],
    "freight_mode": ["freight_mode", "shipping_mode", "mode_of_transport"],
    "waste_type": ["waste_type", "waste_material"],
    "disposal_method": ["disposal_method", "treatment_method", "waste_treatment"],
    "notes": ["notes_context", "notes", "context", "remarks", "memo", "comments"],
}


def _image_to_base64(image: Image.Image) -> str:
    if image.mode in {"RGBA", "P", "LA"}:
        image = image.convert("RGB")
    buffer = BytesIO()
    image.save(buffer, format="JPEG")
    return base64.b64encode(buffer.getvalue()).decode("utf-8")


def render_document(path: str) -> list[str]:
    images: list[str] = []
    if path.lower().endswith(".pdf"):
        with fitz.open(path) as document:
            for index in range(min(len(document), MAX_PDF_PAGES)):
                pixmap = document.load_page(index).get_pixmap(dpi=150, alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                images.append(_image_to_base64(image))
    else:
        with Image.open(path) as image:
            images.append(_image_to_base64(image))
    return images


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _column_mapping(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for key, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_normalize_header(alias) for alias in aliases}
        for index, header in enumerate(headers):
            if header in normalized_aliases:
                mapping[key] = index
                break
    for key, aliases in HEADER_ALIASES.items():
        if key in mapping:
            continue
        for index, header in enumerate(headers):
            if index in mapping.values():
                continue
            if any(len(alias) >= 4 and _normalize_header(alias) in header for alias in aliases):
                mapping[key] = index
                break
    return mapping


def read_spreadsheet(path: str) -> list[dict]:
    rows: list[list[Any]] = []
    if path.lower().endswith(".csv"):
        with open(path, "r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
            rows = [row for row in csv.reader(handle) if any(str(cell).strip() for cell in row)]
    elif path.lower().endswith(".xls"):
        workbook = xlrd.open_workbook(path)
        sheet = workbook.sheet_by_index(0)
        rows = [sheet.row_values(index) for index in range(sheet.nrows) if any(str(cell).strip() for cell in sheet.row_values(index))]
    else:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_name = next((name for name in workbook.sheetnames if any(token in name.lower() for token in ("invoice", "data", "ledger"))), workbook.sheetnames[0])
        rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True) if any(cell not in (None, "") for cell in row)]
    if not rows:
        return []
    header_index = max(
        range(min(15, len(rows))),
        key=lambda idx: sum(any(alias in _normalize_header(cell) for values in HEADER_ALIASES.values() for alias in values) for cell in rows[idx]),
    )
    headers = [_normalize_header(value) for value in rows[header_index]]
    columns = _column_mapping(headers)
    extracted: list[dict] = []
    for raw in rows[header_index + 1:]:
        def value(key: str) -> Any:
            index = columns.get(key)
            return raw[index] if index is not None and index < len(raw) else None
        description = str(value("item_description") or "").strip()
        if not description:
            continue
        facility = str(value("facility") or "").strip()
        reporting_period = value("reporting_period")
        origin = value("origin")
        destination = value("destination")
        distance = value("distance_km")
        extracted.append({
            "invoice_number": str(value("invoice_number") or f"ROW-{len(extracted) + 1:03d}"),
            "date": value("date"),
            "billing_period_text": reporting_period,
            "vendor_name": str(value("vendor_name") or "Corporate Expenditure"),
            "vendor_type": "Supplier/Vendor",
            "service_address": facility,
            "currency": value("currency"),
            "line_items": [{
                "item_description": description,
                "item_description_english": description,
                "item_category_hint": "Other",
                "material_nature": "composite_product",
                "primary_material": "",
                "quantity": value("quantity"),
                "unit": value("unit"),
                "total_cost": value("total_cost"),
                "currency": value("currency"),
                "distance_km": distance,
                "origin": origin,
                "destination": destination,
                "travel_details": {
                    "mode": None, "class": value("travel_class"),
                    "vehicle_type": value("vehicle_type"), "passenger_count": 1,
                    "distance_km": distance, "origin": origin, "destination": destination,
                },
                "freight_details": {
                    "mode": "Road", "distance_km": distance,
                    "origin": origin, "destination": destination,
                },
                "additional_context": value("notes"),
                "confidence_score": None,
                "low_confidence_fields": [],
            }],
            "_spreadsheet_facility_present": bool(facility),
            "_spreadsheet_reporting_period_present": bool(reporting_period),
        })
    return extracted


def _detect_location(invoice: dict, org_context: dict) -> str:
    address = str(invoice.get("service_address") or invoice.get("buyer_address") or "").strip()
    address_lower = address.lower()
    for location in org_context.get("locations", []):
        city = str(location.get("city") or "").strip()
        name = str(location.get("name") or "").strip()
        if (city and city.lower() in address_lower) or (name and name.lower() in address_lower):
            return name or city
    return address or "Unknown"


def _clean_invoice(invoice: dict, org_context: dict) -> tuple[dict, list[dict]]:
    invoice_date, future_date = normalize_date(invoice.get("date"))
    period = normalize_period(
        invoice.get("billing_period_start"),
        invoice.get("billing_period_end"),
        invoice.get("billing_period_text") or invoice.get("billing_period"),
        invoice_date,
    )
    financials = invoice.get("invoice_financials") or {}
    metadata = {
        **invoice,
        "date": invoice_date,
        "billing_period": period,
        "service_address": invoice.get("service_address") or invoice.get("buyer_address"),
        "detected_location": _detect_location(invoice, org_context),
        "currency": normalize_currency(financials.get("currency") or invoice.get("currency")),
        "future_date": future_date,
    }
    if invoice.get("_spreadsheet_facility_present") is not None:
        metadata["vendor_type"] = invoice.get("vendor_type") or "Supplier/Vendor"
        metadata["buyer_name"] = invoice.get("buyer_name") or org_context.get("company_name", "")
        metadata["buyer_address"] = invoice.get("buyer_address") or metadata.get("service_address")
    reconciled = reconcile_invoice_items(invoice)
    invoice_freight = invoice.get("freight_details") or {}
    for item in reconciled:
        if not item.get("freight_details") and invoice_freight:
            item["freight_details"] = invoice_freight
    return metadata, reconciled


async def process_document(
    path: str,
    gateway: OcrLlmGateway,
    org_context: dict,
    enabled_scopes: set[str],
    disabled_scope3_sheets: set[str],
    override_lookup,
) -> list[dict]:
    is_spreadsheet = Path(path).suffix.lower() in {".csv", ".xlsx", ".xls"}
    if is_spreadsheet:
        invoices = read_spreadsheet(path)
    else:
        images = render_document(path)
        response = await gateway.extract_document(EXTRACTION_SYSTEM_PROMPT, EXTRACTION_SCHEMA_PROMPT, images)
        payload = repair_extraction_json(response)
        if isinstance(payload, dict):
            invoices = [payload]
        elif isinstance(payload, list):
            invoices = payload
        else:
            invoices = []
    rows: list[dict] = []
    for invoice in invoices or []:
        if not isinstance(invoice, dict):
            continue
        metadata, items = _clean_invoice(invoice, org_context)
        for item in items:
            if is_spreadsheet:
                item = prepare_spreadsheet_item(
                    item,
                    future_date=metadata.get("future_date", False),
                    invoice_date=metadata.get("date"),
                )
            override = await override_lookup(metadata.get("vendor_name"), item.get("item_description_english") or item.get("item_description"))
            classification = await classify_item(
                gateway,
                item,
                metadata,
                org_context,
                enabled_scopes,
                disabled_scope3_sheets,
                override,
            )
            if is_spreadsheet:
                item = finalize_spreadsheet_item(
                    item,
                    category=classification["ghg_category"],
                    facility_present=bool(invoice.get("_spreadsheet_facility_present")),
                    reporting_period_present=bool(invoice.get("_spreadsheet_reporting_period_present")),
                    classification_needs_review=classification["needs_review"],
                )
            else:
                item = prepare_item(
                    item,
                    future_date=metadata.get("future_date", False),
                    invoice_date=metadata.get("date"),
                )
            confidence = item["confidence_score"]
            freight = item.get("freight_details") or {}
            travel = item.get("travel_details") or {}
            origin = travel.get("origin") or freight.get("origin") or item.get("origin")
            destination = travel.get("destination") or freight.get("destination") or item.get("destination")
            warnings = item.get("quality_warnings") or []
            if "unit" in classification["accounting_rationale"].lower():
                warnings = [warning for warning in warnings if "ASSUMPTION" not in warning]
            rationale = " | ".join([*warnings, classification["accounting_rationale"]]) if warnings else classification["accounting_rationale"]
            row = {
                "invoice_number": metadata.get("invoice_number"),
                "date": metadata.get("date"),
                "billing_period": metadata.get("billing_period"),
                "vendor_name": metadata.get("vendor_name"),
                "vendor_type": metadata.get("vendor_type"),
                "location": metadata.get("detected_location"),
                "item_description": item.get("item_description_english") or item.get("item_description"),
                "item_category_hint": item.get("item_category_hint"),
                "material_nature": item.get("material_nature"),
                "primary_material": item.get("primary_material"),
                "hsn_sac_code": item.get("hsn_sac_code"),
                "unit_price": item.get("unit_price"),
                "waste_details": item.get("waste_details"),
                "travel_details": travel,
                "freight_details": freight,
                "scope": classification["ghg_scope"],
                "category": classification["ghg_category"],
                "category_key": classification["category_key"],
                "category_code": classification["category_code"],
                "subcategory": classification["ghg_subcategory"],
                "fuel_name": classification["ghg_subcategory"],
                "quantity": item.get("quantity"),
                "unit": normalize_unit(item.get("unit"), classification["ghg_category"]),
                "distance_km": item.get("distance_km"),
                "origin": origin,
                "destination": destination,
                "cost": item.get("total_cost"),
                "currency": normalize_currency(item.get("currency") or metadata.get("currency")),
                "ef_method": classification["ef_method"],
                "ef_database": classification["ef_database"],
                "ef_lookup_key": classification["ef_lookup_key"],
                "naics_code": classification["naics_code"],
                "naics_label": classification["naics_label"],
                "accounting_rationale": rationale,
                "confidence_score": confidence,
                "classification_confidence_score": classification["confidence_score"],
                "low_confidence_fields": item.get("low_confidence_fields", []),
                "needs_review": item.get("needs_review") if is_spreadsheet else (item.get("missing_values") or classification["needs_review"] or confidence < 70 or bool(item.get("low_confidence_fields"))),
                "missing_values": item.get("missing_values", False),
                "is_auto_generated": False,
                "classification_source": classification["classification_source"],
                "auto_generate_cat3": classification["auto_generate_cat3"],
            }
            rows.append(row)
            rows.extend(build_companion_rows(row))
    return rows
