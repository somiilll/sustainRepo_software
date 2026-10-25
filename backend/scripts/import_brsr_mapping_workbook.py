"""Import validated BRSR mapping choices and export a revised review workbook."""
import asyncio
import json
import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from modules.brsr_report.mapping_excel import build_kpi_brsr_mapping_excel, logical_answer_path
from shared.database.mongo import db


WORKBOOK_PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/app/uploads/user_kpi_definitions_brsr_mappings.xlsx")
EXPORT_PATH = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("/app/exports/kpi_definitions_brsr_mappings_revised.xlsx")
BACKUP_DIR = Path("/app/backups")


def _normalise(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _complaint_destination(name: str, code: str) -> Optional[Tuple[str, str]]:
    text = _normalise(f"{name} {code}")
    pending = " pending " in f" {text} " or "pending resolution" in text
    if "consumer complaint" in text:
        rows = {
            "data privacy": "data_privacy", "advertising": "advertising", "cyber security": "cyber_security",
            "delivery of essential services": "delivery_essential", "restrictive unfair trade": "unfair_trade",
            "restrictive trade": "restrictive_trade", "unfair trade": "unfair_trade", "other": "other",
        }
        row = next((row_key for token, row_key in rows.items() if token in text), None)
        if row:
            return "p9_consumer_complaints", f"rows.{row}.{'pending' if pending else 'received'}"
    if "posh" in text:
        row = "posh_upheld" if "upheld" in text else "total_reported"
        return "p5_sexual_harassment_complaints", f"rows.{row}"
    if "internal complaint" in text:
        row_map = {
            "working conditions": ("p3_complaints_employees_workers", "working_conditions"),
            "health safety": ("p3_complaints_employees_workers", "health_safety"),
            "sexual harassment": ("p5_hr_complaints", "sexual_harassment"),
            "discrimination": ("p5_hr_complaints", "discrimination"),
            "child labour": ("p5_hr_complaints", "child_labour"),
            "forced labour": ("p5_hr_complaints", "forced_labour"),
            "wages": ("p5_hr_complaints", "wages"),
            "other human rights": ("p5_hr_complaints", "other_hr"),
        }
        matched = next((destination for token, destination in row_map.items() if token in text), None)
        if matched:
            question_key, row = matched
            return question_key, f"rows.{row}.{'pending' if pending else 'filed'}"
    return None


def _section_a_destination(name: str, code: str, reference: str) -> Optional[Tuple[str, str]]:
    text = _normalise(f"{name} {code}")
    reference = str(reference or "")
    if reference.startswith("A-21"):
        if "board of directors" in text or "bod " in f"{text} ":
            row = "board_of_directors"
        elif "key management personnel" in text or "kmp " in f"{text} ":
            row = "key_management_personnel"
        else:
            return None
        if "female" in text:
            field = "number_of_females"
        elif "male" in text:
            return None
        elif "total" in text:
            field = "total"
        else:
            return None
        return "brsr_a_women_representation", f"rows.{row}.{field}"

    if reference.startswith("A-20"):
        differently_abled = "differently abled" in text or " div da " in f" {text} "
        question_key = "brsr_a_differently_abled" if differently_abled else "brsr_a_employees_workers"
        category = "workers" if "worker" in text else "employees" if "employee" in text else None
        if not category:
            return None
        if "other than permanent" in text or "other permanent" in text:
            employment_type = "other_than_permanent"
        elif "permanent" in text:
            employment_type = "permanent"
        else:
            return None
        if "female" in text:
            field = "female"
        elif "male" in text:
            field = "male"
        elif "total" in text:
            field = "total"
        else:
            return None
        return question_key, f"{category}.{employment_type}.{field}"
    return None


def _direct_destination(name: str, code: str, reference: str) -> Optional[Tuple[str, str]]:
    text = _normalise(f"{name} {code}")
    complaint = _complaint_destination(name, code)
    if complaint:
        return complaint
    section_a = _section_a_destination(name, code, reference)
    if section_a:
        return section_a
    if "accounts payable days" in text:
        return "p1_accounts_payables", "rows.days_payables"
    if "data breach incidents total" in text or "data breach incidents" in text and "percentage" not in text:
        return "p9_data_breaches_count", "value"
    if "data breach" in text and ("percentage" in text or "customer pii" in text):
        return "p9_data_breaches_pii_pct", "value"
    return None


def _refine_selected_path(name: str, code: str, question_key: str, answer_path: str) -> Tuple[str, str]:
    answer_path = logical_answer_path(answer_path) or ""
    if question_key in {"p9_consumer_complaints", "p5_hr_complaints", "p3_complaints_employees_workers",
                        "p5_sexual_harassment_complaints"} and answer_path and answer_path.count(".") == 1:
        inferred = _complaint_destination(name, code)
        if inferred and inferred[0] == question_key:
            return inferred
    return question_key, answer_path


def _energy_destination_is_safe(name: str, path: str) -> bool:
    text = _normalise(name)
    expected = None
    if "grid electricity" in text:
        expected = "rows.nonrenew_electricity"
    elif "renewable electricity" in text:
        expected = "rows.renew_electricity"
    return expected == path if expected else False


def _waste_destination_is_safe(name: str, path: str) -> bool:
    text = _normalise(name)
    if "waste disposed incineration" in text:
        return path == "rows.disposed_incineration"
    if "waste disposed landfilling" in text:
        return path == "rows.disposed_landfilling"
    if "waste disposed other disposal" in text:
        return path == "rows.disposed_other"
    if "waste disposed total" in text:
        return path == "rows.disposed_total"
    if "waste recovered recycled" in text:
        return path == "rows.recovered_recycled"
    if "waste recovered re used" in text:
        return path == "rows.recovered_reused"
    if "waste recovered other recovery" in text:
        return path == "rows.recovered_other"
    if "waste recovered total" in text:
        return path == "rows.recovered_total"
    generated_rows = {
        "plastic generated": "rows.plastic_waste", "e waste generated": "rows.e_waste",
        "biomedical generated": "rows.bio_medical", "construction demolition generated": "rows.construction",
        "battery generated": "rows.battery", "radioactive generated": "rows.radioactive",
        "other hazardous waste generated": "rows.other_hazardous",
        "other non hazardous waste generated": "rows.other_non_hazardous",
    }
    expected = next((row for token, row in generated_rows.items() if token in text), None)
    if expected:
        return path == expected
    if "waste generated total" in text:
        return path == "rows.total_generated"
    return False


def _period_metadata(config: Dict[str, Any], answer_path: str) -> Tuple[str, Dict[str, str]]:
    if config.get("section") in {"section_a", "section_b"} or config.get("type") in {
        "text", "number", "email", "url", "radio", "currency", "textarea"
    }:
        return "response_period_document", {}
    return "reporting_year_filter", {}


async def _apply_mapping(
    kpi: Dict[str, Any], config: Dict[str, Any], question_key: str, answer_path: str,
    existing_mappings: List[Dict[str, Any]], used_mapping_ids: set, source: str, now: str,
) -> Dict[str, Any]:
    identical = next((mapping for mapping in existing_mappings
                      if mapping.get("question_key") == question_key and mapping.get("answer_path") == answer_path), None)
    target = identical or next((mapping for mapping in existing_mappings
                                if mapping.get("id") not in used_mapping_ids and
                                mapping.get("mapping_status") != "exact_destination"), None)
    strategy, period_columns = _period_metadata(config, answer_path)
    fields = {
        "kpi_id": kpi["id"], "source_kpi_ids": [kpi["id"]], "source_metric_code": kpi.get("metric_code"),
        "source_metric_name": kpi.get("metric_name"), "framework": "BRSR",
        "reference": kpi.get("brsr_reference") or (target or {}).get("reference"),
        "question_key": question_key, "answer_path": answer_path, "destination_label": config.get("question"),
        "aggregation_rule": (target or {}).get("aggregation_rule") or kpi.get("aggregation_type") or "sum",
        "formula_definition": (target or {}).get("formula_definition"),
        "destination_unit": (target or {}).get("destination_unit"),
        "unit_conversion_rule": (target or {}).get("unit_conversion_rule"),
        "mapping_status": "exact_destination", "is_active": True,
        "period_column_strategy": strategy, "period_columns": period_columns,
        "source_category": kpi.get("category_name"), "source_subcategory": kpi.get("subcategory"),
        "source_sub_subcategory": kpi.get("sub_subcategory"),
        "source_tag": " → ".join(filter(None, [kpi.get("category_name"), kpi.get("subcategory")])),
        "review_source": source, "reviewed_at": now, "updated_at": now,
    }
    if target:
        await db.kpi_framework_mappings.update_one({"id": target["id"]}, {"$set": fields})
        mapping_id = target["id"]
        action = "updated"
    else:
        mapping_id = str(uuid.uuid4())
        await db.kpi_framework_mappings.insert_one({"id": mapping_id, "created_at": now, **fields})
        existing_mappings.append({"id": mapping_id, **fields})
        action = "created"
    used_mapping_ids.add(mapping_id)
    return {"action": action, "mapping_id": mapping_id, "kpi_code": kpi.get("metric_code"),
            "kpi_name": kpi.get("metric_name"), "question_key": question_key, "answer_path": answer_path,
            "source": source}


async def _apply_energy_group(
    definitions: List[Dict[str, Any]], config: Dict[str, Any], renewable: bool, now: str
) -> Optional[Dict[str, Any]]:
    group = []
    for definition in definitions:
        text = _normalise(definition.get("metric_name"))
        if not any(source in text for source in ("cooling", "heating", "steam")):
            continue
        if renewable and "renewable energy consumption" in text and "non renewable" not in text:
            group.append(definition)
        if not renewable and "non renewable energy consumption" in text:
            group.append(definition)
    if not group:
        return None
    path = "rows.renew_other" if renewable else "rows.nonrenew_other"
    ids = sorted(definition["id"] for definition in group)
    existing = await db.kpi_framework_mappings.find_one({
        "framework": "BRSR", "question_key": "p6_energy_consumption", "answer_path": path,
        "mapping_status": "exact_destination", "source_kpi_ids": {"$all": ids, "$size": len(ids)},
    }, {"_id": 0})
    fields = {
        "kpi_id": ids[0], "source_kpi_ids": ids,
        "source_metric_code": "GROUP_RENEWABLE_OTHER_ENERGY" if renewable else "GROUP_NONRENEWABLE_OTHER_ENERGY",
        "source_metric_name": "Renewable other-source energy" if renewable else "Non-renewable other-source energy",
        "framework": "BRSR", "reference": "P6-E1 (C)" if renewable else "P6-E1 (F)",
        "question_key": "p6_energy_consumption", "answer_path": path,
        "destination_label": config.get("question"), "aggregation_rule": "sum", "formula_definition": None,
        "destination_unit": None, "unit_conversion_rule": "normalize_to_destination_unit",
        "mapping_status": "exact_destination", "is_active": True,
        "period_column_strategy": "reporting_period_relative",
        "period_columns": {"selected_reporting_year": "current_fy", "previous_reporting_year": "previous_fy",
                           "prior_to_previous_reporting_year": "prior_to_previous_fy"},
        "source_category": "Energy", "source_subcategory": "Multiple energy sources",
        "source_sub_subcategory": None, "source_tag": "Energy → Cooling / Heating / Steam",
        "review_source": "high_confidence_grouping", "reviewed_at": now, "updated_at": now,
    }
    if existing:
        await db.kpi_framework_mappings.update_one({"id": existing["id"]}, {"$set": fields})
        mapping_id = existing["id"]
        action = "updated"
    else:
        mapping_id = str(uuid.uuid4())
        await db.kpi_framework_mappings.insert_one({"id": mapping_id, "created_at": now, **fields})
        action = "created"
    await db.kpi_framework_mappings.update_many(
        {"framework": "BRSR", "id": {"$ne": mapping_id}, "kpi_id": {"$in": ids},
         "mapping_status": {"$ne": "exact_destination"}},
        {"$set": {"is_active": False, "mapping_status": "superseded_by_group_mapping",
                  "superseded_by_mapping_id": mapping_id, "updated_at": now}},
    )
    return {"action": action, "mapping_id": mapping_id, "kpi_code": fields["source_metric_code"],
            "kpi_name": fields["source_metric_name"], "question_key": fields["question_key"],
            "answer_path": path, "source": "high_confidence_grouping"}


async def run() -> Dict[str, Any]:
    now = datetime.now(timezone.utc).isoformat()
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    backup_path = BACKUP_DIR / f"kpi_framework_mappings_before_workbook_import_{now.replace(':', '').replace('+', '_')}.json"
    existing_all = await db.kpi_framework_mappings.find({"framework": "BRSR"}, {"_id": 0}).to_list(10000)
    backup_path.write_text(json.dumps(existing_all, default=str, indent=2), encoding="utf-8")

    configs = await db.esg_question_configs.find(
        {"frameworks": "BRSR", "hidden": {"$ne": True}}, {"_id": 0}
    ).to_list(2000)
    config_by_key = {config["question_key"]: config for config in configs}
    allowed_paths = {key: set(config.get("answer_paths") or []) for key, config in config_by_key.items()}
    definitions = await db.esg_kpi_definitions.find({"status": {"$ne": "archived"}}, {"_id": 0}).to_list(10000)
    definitions_by_code = {definition.get("metric_code"): definition for definition in definitions if definition.get("metric_code")}
    mappings_by_kpi: Dict[str, List[Dict[str, Any]]] = {}
    for mapping in existing_all:
        mappings_by_kpi.setdefault(mapping.get("kpi_id"), []).append(mapping)

    workbook = load_workbook(WORKBOOK_PATH, data_only=False)
    worksheet = workbook["BRSR KPI Mappings"]
    applied: List[Dict[str, Any]] = []
    review: List[Dict[str, Any]] = []
    used_mapping_ids: set = set()
    selected_codes = set()

    for row_number in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row_number, 1).value
        code = worksheet.cell(row_number, 2).value
        reference = worksheet.cell(row_number, 6).value
        question_key = worksheet.cell(row_number, 7).value
        answer_path = worksheet.cell(row_number, 8).value
        if not question_key and not answer_path:
            continue
        selected_codes.add(code)
        kpi = definitions_by_code.get(code)
        if not kpi:
            review.append({"row": row_number, "kpi_name": name, "kpi_code": code,
                           "question_key": question_key, "answer_path": answer_path,
                           "reason": "KPI code is not an active KPI Definition"})
            continue
        question_key, answer_path = _refine_selected_path(name, code, str(question_key or ""), str(answer_path or ""))
        config = config_by_key.get(question_key)
        if not config:
            review.append({"row": row_number, "kpi_name": name, "kpi_code": code,
                           "question_key": question_key, "answer_path": answer_path,
                           "reason": "Question Key is not a configured active BRSR question"})
            continue
        if answer_path not in allowed_paths.get(question_key, set()):
            inferred = _direct_destination(name, code, reference)
            if inferred and inferred[0] == question_key and inferred[1] in allowed_paths.get(question_key, set()):
                answer_path = inferred[1]
            else:
                review.append({"row": row_number, "kpi_name": name, "kpi_code": code,
                               "question_key": question_key, "answer_path": answer_path,
                               "reason": "Answer Path is not valid for the selected BRSR question"})
                continue
        if question_key == "p6_energy_consumption" and not _energy_destination_is_safe(name, answer_path):
            review.append({"row": row_number, "kpi_name": name, "kpi_code": code,
                           "question_key": question_key, "answer_path": answer_path,
                           "reason": "This KPI does not independently equal the selected Energy total row"})
            continue
        if question_key == "p6_waste_management_details" and not _waste_destination_is_safe(name, answer_path):
            review.append({"row": row_number, "kpi_name": name, "kpi_code": code,
                           "question_key": question_key, "answer_path": answer_path,
                           "reason": "This component KPI does not independently equal the selected Waste total row"})
            continue
        applied.append(await _apply_mapping(
            kpi, config, question_key, answer_path, mappings_by_kpi.setdefault(kpi["id"], []),
            used_mapping_ids, "user_workbook", now,
        ))

    for kpi in definitions:
        destination = _direct_destination(kpi.get("metric_name"), kpi.get("metric_code"), kpi.get("brsr_reference"))
        if not destination:
            continue
        question_key, answer_path = destination
        config = config_by_key.get(question_key)
        if not config or answer_path not in allowed_paths.get(question_key, set()):
            continue
        existing = mappings_by_kpi.setdefault(kpi["id"], [])
        if any(mapping.get("mapping_status") == "exact_destination" and
               mapping.get("question_key") == question_key and mapping.get("answer_path") == answer_path
               for mapping in existing):
            continue
        applied.append(await _apply_mapping(
            kpi, config, question_key, answer_path, existing, used_mapping_ids,
            "high_confidence_inference", now,
        ))

    energy_config = config_by_key.get("p6_energy_consumption")
    if energy_config:
        for renewable in (True, False):
            grouped = await _apply_energy_group(definitions, energy_config, renewable, now)
            if grouped:
                applied.append(grouped)

    workbook_bytes = await build_kpi_brsr_mapping_excel()
    EXPORT_PATH.write_bytes(workbook_bytes)
    revised = load_workbook(EXPORT_PATH)
    if "Workbook Import Review" in revised.sheetnames:
        del revised["Workbook Import Review"]
    import_review = revised.create_sheet("Workbook Import Review", 1)
    headers = ["Original Row", "KPI Name", "KPI Code", "Selected Question Key", "Selected Answer Path", "Review Reason"]
    import_review.append(headers)
    for item in review:
        import_review.append([item.get("row"), item.get("kpi_name"), item.get("kpi_code"),
                              item.get("question_key"), item.get("answer_path"), item.get("reason")])
    for cell in import_review[1]:
        cell.fill = PatternFill("solid", fgColor="92400E")
        cell.font = Font(color="FFFFFF", bold=True)
    import_review.freeze_panes = "A2"
    for column, width in zip("ABCDEF", [14, 42, 54, 34, 44, 58]):
        import_review.column_dimensions[column].width = width

    applied_sheet = revised.create_sheet("Applied Mapping Summary", 1)
    applied_sheet.append(["Action", "KPI Name", "KPI Code", "Question Key", "Answer Path", "Source"])
    for item in applied:
        applied_sheet.append([item["action"], item["kpi_name"], item["kpi_code"], item["question_key"],
                              item["answer_path"], item["source"]])
    for cell in applied_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="14532D")
        cell.font = Font(color="FFFFFF", bold=True)
    applied_sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEF", [14, 42, 54, 34, 44, 28]):
        applied_sheet.column_dimensions[column].width = width
    revised.save(EXPORT_PATH)

    log_path = EXPORT_PATH.with_suffix(".json")
    result = {"applied_count": len(applied), "review_count": len(review), "backup": str(backup_path),
              "export": str(EXPORT_PATH), "log": str(log_path)}
    log_path.write_text(json.dumps({"result": result, "applied": applied, "review": review}, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    print(asyncio.run(run()))